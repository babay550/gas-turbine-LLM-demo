"""数据字典 API — 参数定义、基准值配置、对标指标。

从 v1.8 起所有数据从 SQLite 数据库读写，首次启动自动 seed。
支持 XLSX 模板下载/导入/导出批量管理参数。
"""

import io
import json
import logging
from datetime import datetime

from fastapi import APIRouter, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from src.db.engine import get_session
from src.db.models import Parameter, BaselineConfig, BenchmarkIndicator, LossVariableConfig, WaterfallConfig

logger = logging.getLogger(__name__)
router = APIRouter()

# ──────────── 字典 seed 数据（从生产 DB 反向同步，首次启动 seed 用） ────────────
# engine.py 的 _seed_if_empty() 读取；rebuild 镜像后新环境开箱即用

PARAMETERS = [
    {"id": "P001", "name": "压气机进口温度 T1", "unit": "°C", "subsystem": "压气机", "location": "压气机进口", "normal_range": [10, 40], "source": "时序数据库", "update_freq": "1s"},
    {"id": "P002", "name": "压气机进口压力 P1", "unit": "MPa", "subsystem": "压气机", "location": "压气机进口", "normal_range": [0.09, 0.11], "source": "时序数据库", "update_freq": "1s"},
    {"id": "P003", "name": "压气机出口温度 T2", "unit": "°C", "subsystem": "压气机", "location": "压气机出口", "normal_range": [380, 460], "source": "时序数据库", "update_freq": "1s"},
    {"id": "P004", "name": "压气机出口压力 P2", "unit": "MPa", "subsystem": "压气机", "location": "压气机出口", "normal_range": [1.2, 1.9], "source": "时序数据库", "update_freq": "1s"},
    {"id": "P005", "name": "空气流量", "unit": "kg/s", "subsystem": "压气机", "location": "压气机出口", "normal_range": [480, 720], "source": "时序数据库", "update_freq": "1s"},
    {"id": "P006", "name": "透平进口温度 T3", "unit": "°C", "subsystem": "燃烧室", "location": "燃烧室出口", "normal_range": [1050, 1350], "source": "时序数据库", "update_freq": "1s"},
    {"id": "P007", "name": "燃烧室出口压力 P3", "unit": "MPa", "subsystem": "燃烧室", "location": "燃烧室出口", "normal_range": [1.1, 1.8], "source": "时序数据库", "update_freq": "1s"},
    {"id": "P008", "name": "天然气瞬时流量", "unit": "万Nm³/h", "subsystem": "燃烧室", "location": "燃料入口", "normal_range": [4, 7], "source": "时序数据库", "update_freq": "1s"},
    {"id": "P009", "name": "天然气累计流量", "unit": "万Nm³", "subsystem": "燃烧室", "location": "燃料入口", "normal_range": [0, 99999], "source": "时序数据库", "update_freq": "1h"},
    {"id": "P010", "name": "燃气温度", "unit": "°C", "subsystem": "燃烧室", "location": "燃料入口", "normal_range": [0, 40], "source": "时序数据库", "update_freq": "1s"},
    {"id": "P011", "name": "燃气压力", "unit": "MPa", "subsystem": "燃烧室", "location": "燃料入口", "normal_range": [1.5, 4], "source": "时序数据库", "update_freq": "1s"},
    {"id": "P012", "name": "天然气低位热值", "unit": "MJ/Nm³", "subsystem": "燃烧室", "location": "化验室", "normal_range": [31, 38], "source": "定期化验", "update_freq": "1天"},
    {"id": "P013", "name": "排烟含氧量", "unit": "%", "subsystem": "燃烧室", "location": "烟囱", "normal_range": [10, 18], "source": "时序数据库", "update_freq": "1s"},
    {"id": "P014", "name": "CO排放", "unit": "ppm", "subsystem": "燃烧室", "location": "烟囱", "normal_range": [0, 20], "source": "时序数据库", "update_freq": "1s"},
    {"id": "P015", "name": "未燃碳氢", "unit": "ppm", "subsystem": "燃烧室", "location": "烟囱", "normal_range": [0, 10], "source": "时序数据库", "update_freq": "1s"},
    {"id": "P016", "name": "透平出口温度 T4", "unit": "°C", "subsystem": "透平", "location": "透平出口", "normal_range": [500, 600], "source": "时序数据库", "update_freq": "1s"},
    {"id": "P017", "name": "透平出口压力 P4", "unit": "MPa", "subsystem": "透平", "location": "透平出口", "normal_range": [0.09, 0.11], "source": "时序数据库", "update_freq": "1s"},
    {"id": "P018", "name": "发电机有功功率", "unit": "MW", "subsystem": "发电机", "location": "发电机出口", "normal_range": [160, 230], "source": "时序数据库", "update_freq": "1s"},
    {"id": "P019", "name": "累计发电量", "unit": "MWh", "subsystem": "发电机", "location": "发电机出口", "normal_range": [0, 999999], "source": "时序数据库", "update_freq": "1h"},
    {"id": "P020", "name": "主蒸汽流量", "unit": "t/h", "subsystem": "汽机", "location": "余热锅炉出口", "normal_range": [50, 100], "source": "时序数据库", "update_freq": "1s"},
    {"id": "P021", "name": "主蒸汽温度", "unit": "°C", "subsystem": "汽机", "location": "余热锅炉出口", "normal_range": [470, 550], "source": "时序数据库", "update_freq": "1s"},
    {"id": "P022", "name": "主蒸汽压力", "unit": "MPa", "subsystem": "汽机", "location": "余热锅炉出口", "normal_range": [3, 5.5], "source": "时序数据库", "update_freq": "1s"},
    {"id": "P023", "name": "汽机发电功率", "unit": "MW", "subsystem": "汽机", "location": "汽轮发电机", "normal_range": [70, 130], "source": "时序数据库", "update_freq": "1s"},
    {"id": "P024", "name": "HRB进口烟温", "unit": "°C", "subsystem": "余热锅炉", "location": "余热锅炉进口", "normal_range": [500, 600], "source": "时序数据库", "update_freq": "1s"},
    {"id": "P025", "name": "HRB出口烟温", "unit": "°C", "subsystem": "余热锅炉", "location": "余热锅炉出口", "normal_range": [70, 130], "source": "时序数据库", "update_freq": "1s"},
    {"id": "P026", "name": "HRB蒸汽温度", "unit": "°C", "subsystem": "余热锅炉", "location": "余热锅炉", "normal_range": [470, 540], "source": "时序数据库", "update_freq": "1s"},
    {"id": "P027", "name": "HRB蒸汽压力", "unit": "MPa", "subsystem": "余热锅炉", "location": "余热锅炉", "normal_range": [2.5, 5], "source": "时序数据库", "update_freq": "1s"},
    {"id": "P028", "name": "振动_轴向", "unit": "mm/s", "subsystem": "辅助", "location": "轴承座", "normal_range": [0, 7], "source": "时序数据库", "update_freq": "1s"},
    {"id": "P029", "name": "振动_垂向", "unit": "mm/s", "subsystem": "辅助", "location": "轴承座", "normal_range": [0, 7], "source": "时序数据库", "update_freq": "1s"},
    {"id": "P030", "name": "燃机效率_value", "unit": "%", "subsystem": "其他", "normal_range": [34, 40], "source": "时序数据库", "update_freq": "1min"},
    {"id": "P031", "name": "燃机效率_referValue", "unit": "%", "subsystem": "其他", "normal_range": [34, 40], "source": "时序数据库", "update_freq": "1min", "description": "TTDM"},
    {"id": "P032", "name": "燃机效率_coalLossValue", "unit": "g/kWh", "subsystem": "其他", "normal_range": [-3, 3], "source": "时序数据库", "update_freq": "1min"},
    {"id": "P033", "name": "燃机效率_optimalLossValue", "unit": "g/kWh", "subsystem": "其他", "normal_range": [-1, 4], "source": "时序数据库", "update_freq": "1min"},
    {"id": "P034", "name": "余热锅炉效率_value", "unit": "%", "subsystem": "余热锅炉", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P035", "name": "余热锅炉效率_referValue", "unit": "%", "subsystem": "余热锅炉", "source": "时序数据库", "update_freq": "1min", "description": "TTDM"},
    {"id": "P036", "name": "余热锅炉效率_coalValue", "unit": "g/kWh", "subsystem": "余热锅炉", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P037", "name": "余热锅炉效率_coalLossValue", "unit": "g/kWh", "subsystem": "余热锅炉", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P038", "name": "余热锅炉效率_optimalLossValue", "unit": "g/kWh", "subsystem": "余热锅炉", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P039", "name": "汽机效率_value", "unit": "%", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P040", "name": "汽机效率_referValue", "unit": "%", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P041", "name": "汽机效率_coalValue", "unit": "g/kWh", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P042", "name": "汽机效率_coalLossValue", "unit": "g/kWh", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P043", "name": "汽机效率_optimalLossValue", "unit": "g/kWh", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P044", "name": "能耗偏差_value", "unit": "g/kWh", "subsystem": "机组", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P045", "name": "能耗偏差_coalLossValue", "unit": "g/kWh", "subsystem": "机组", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P046", "name": "能耗偏差_optimalLossValue", "unit": "g/kWh", "subsystem": "机组", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P047", "name": "压气机效率_value", "unit": "%", "subsystem": "压气机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P048", "name": "压气机效率_referValue", "unit": "%", "subsystem": "压气机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P049", "name": "压气机效率_coalValue", "unit": "g/kWh", "subsystem": "压气机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P050", "name": "压气机效率_coalLossValue", "unit": "g/kWh", "subsystem": "压气机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P051", "name": "压气机效率_optimalLossValue", "unit": "g/kWh", "subsystem": "压气机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P052", "name": "压气机进气压损_value", "unit": "KPa", "subsystem": "压气机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P053", "name": "压气机进气压损_referValue", "unit": "KPa", "subsystem": "压气机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P054", "name": "压气机进气压损_coalValue", "unit": "g/kWh", "subsystem": "压气机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P055", "name": "压气机进气压损_coalLossValue", "unit": "g/kWh", "subsystem": "压气机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P056", "name": "压气机进气压损_optimalLossValue", "unit": "g/kWh", "subsystem": "压气机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P057", "name": "燃机排气压力_value", "unit": "KPa", "subsystem": "燃机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P058", "name": "燃机排气压力_referValue", "unit": "KPa", "subsystem": "燃机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P059", "name": "燃机排气压力_coalValue", "unit": "g/kWh", "subsystem": "燃机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P060", "name": "燃机排气压力_coalLossValue", "unit": "g/kWh", "subsystem": "燃机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P061", "name": "燃机排气压力_optimalLossValue", "unit": "g/kWh", "subsystem": "燃机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P062", "name": "环境温度_value", "unit": "°C", "subsystem": "环境", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P063", "name": "环境温度_referValue", "unit": "°C", "subsystem": "环境", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P064", "name": "环境温度_coalValue", "unit": "g/kWh", "subsystem": "环境", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P065", "name": "环境温度_coalLossValue", "unit": "g/kWh", "subsystem": "环境", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P066", "name": "环境温度_optimalLossValue", "unit": "g/kWh", "subsystem": "环境", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P067", "name": "大气压力_value", "unit": "KPa", "subsystem": "环境", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P068", "name": "大气压力_referValue", "unit": "KPa", "subsystem": "环境", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P069", "name": "大气压力_coalValue", "unit": "g/kWh", "subsystem": "环境", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P070", "name": "大气压力_coalLossValue", "unit": "g/kWh", "subsystem": "环境", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P071", "name": "大气压力_optimalLossValue", "unit": "g/kWh", "subsystem": "环境", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P072", "name": "大气相对湿度_value", "unit": "%", "subsystem": "环境", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P073", "name": "大气相对湿度_referValue", "unit": "%", "subsystem": "环境", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P074", "name": "大气相对湿度_coalLossValue", "unit": "g/kWh", "subsystem": "环境", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P075", "name": "大气相对湿度_optimalLossValue", "unit": "g/kWh", "subsystem": "环境", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P076", "name": "余热锅炉烟气压损_value", "unit": "KPa", "subsystem": "余热锅炉", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P077", "name": "余热锅炉烟气压损_referValue", "unit": "KPa", "subsystem": "余热锅炉", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P078", "name": "余热锅炉烟气压损_coalLossValue", "unit": "g/kWh", "subsystem": "余热锅炉", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P079", "name": "余热锅炉烟气压损_optimalLossValue", "unit": "g/kWh", "subsystem": "余热锅炉", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P080", "name": "余热锅炉排烟温度_value", "unit": "°C", "subsystem": "余热锅炉", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P081", "name": "余热锅炉排烟温度_referValue", "unit": "°C", "subsystem": "余热锅炉", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P082", "name": "余热锅炉排烟温度_coalLossValue", "unit": "g/kWh", "subsystem": "余热锅炉", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P083", "name": "余热锅炉排烟温度_optimalLossValue", "unit": "g/kWh", "subsystem": "余热锅炉", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P084", "name": "汽机高压主蒸汽压力_value", "unit": "KPa", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P085", "name": "汽机高压主蒸汽压力_referValue", "unit": "KPa", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P086", "name": "汽机高压主蒸汽压力_coalLossValue", "unit": "g/kWh", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P087", "name": "汽机高压主蒸汽压力_optimalLossValue", "unit": "g/kWh", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P088", "name": "汽机高压主蒸汽温度_value", "unit": "°C", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P089", "name": "汽机高压主蒸汽温度_referValue", "unit": "°C", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P090", "name": "汽机高压主蒸汽温度_coalLossValue", "unit": "g/kWh", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P091", "name": "汽机高压主蒸汽温度_optimalLossValue", "unit": "g/kWh", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P092", "name": "汽机低压主蒸汽温度_value", "unit": "°C", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P093", "name": "汽机低压主蒸汽温度_referValue", "unit": "°C", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P094", "name": "汽机低压主蒸汽温度_coalLossValue", "unit": "g/kWh", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P095", "name": "汽机低压主蒸汽温度_optimalLossValue", "unit": "g/kWh", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P096", "name": "凝汽器压力_value", "unit": "KPa", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P097", "name": "凝汽器压力_referValue", "unit": "KPa", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P098", "name": "凝汽器压力_coalLossValue", "unit": "g/kWh", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P099", "name": "凝汽器压力_optimalLossValue", "unit": "g/kWh", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P100", "name": "过冷度_value", "unit": "°C", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P101", "name": "过冷度_referValue", "unit": "°C", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P102", "name": "过冷度_coalLossValue", "unit": "g/kWh", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P103", "name": "过冷度_optimalLossValue", "unit": "g/kWh", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P104", "name": "高压缸效率_value", "unit": "%", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P105", "name": "高压缸效率_referValue", "unit": "%", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P106", "name": "高压缸效率_coalLossValue", "unit": "g/kWh", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P107", "name": "高压缸效率_optimalLossValue", "unit": "g/kWh", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P108", "name": "再热蒸汽压损_value", "unit": "KPa", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P109", "name": "再热蒸汽压损_referValue", "unit": "KPa", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P110", "name": "再热蒸汽压损_coalLossValue", "unit": "g/kWh", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P111", "name": "再热蒸汽压损_optimalLossValue", "unit": "g/kWh", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P112", "name": "再热蒸汽温度_value", "unit": "°C", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P113", "name": "再热蒸汽温度_referValue", "unit": "°C", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P114", "name": "再热蒸汽温度_coalLossValue", "unit": "g/kWh", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P115", "name": "再热蒸汽温度_optimalLossValue", "unit": "g/kWh", "subsystem": "汽机", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P116", "name": "IGV开度_value", "unit": "%", "subsystem": "机组", "source": "时序数据库", "update_freq": "1min"},
    {"id": "P117", "name": "负荷率_value", "unit": "%", "subsystem": "机组", "source": "时序数据库", "update_freq": "1min"},
]

BASELINE_CONFIGS = [
    {"parameter_id": "P003", "parameter_name": "压气机出口温度 T2", "model_type": "多元回归", "features": ["负荷率", "进口温度 T1", "进口压力 P1"], "accuracy": 0.98},
    {"parameter_id": "P016", "parameter_name": "透平出口温度 T4", "model_type": "多元回归", "features": ["负荷率", "环境温度", "透平进口温度 T3"], "accuracy": 0.97},
    {"parameter_id": "P008", "parameter_name": "天然气瞬时流量", "model_type": "设计值插值", "features": ["负荷指令", "发电功率"], "accuracy": 0.99},
    {"parameter_id": "P018", "parameter_name": "发电机有功功率", "model_type": "热力模型", "features": ["天然气流量", "热值", "压气机效率"], "accuracy": 0.96},
]

BENCHMARK_INDICATORS = [
    {"id": "B001", "name": "厂用电率", "unit": "%", "source": "对标分析模型", "peer_avg": 4.8},
    {"id": "B002", "name": "利用小时", "unit": "h", "source": "对标分析模型", "peer_avg": 4200},
    {"id": "B003", "name": "平均发电负荷", "unit": "MW", "source": "对标分析模型", "peer_avg": 200},
    {"id": "B004", "name": "负荷率", "unit": "%", "source": "对标分析模型", "peer_avg": 90},
    {"id": "B005", "name": "发电计划完成率", "unit": "%", "source": "对标分析模型", "peer_avg": 95},
    {"id": "B006", "name": "热电比", "unit": "", "source": "对标分析模型", "peer_avg": 0.85},
    {"id": "B007", "name": "售电收入", "unit": "万元", "source": "对标分析模型", "peer_avg": 12000},
]


# ──────────── 请求模型 ────────────

class ParameterCreate(BaseModel):
    name: str
    unit: str
    subsystem: str = "其他"
    location: str = ""
    normal_min: float | None = None
    normal_max: float | None = None
    source: str = "时序数据库"
    update_freq: str = "1s"
    description: str = ""


class ParameterUpdate(BaseModel):
    name: str | None = None
    unit: str | None = None
    subsystem: str | None = None
    location: str | None = None
    normal_min: float | None = None
    normal_max: float | None = None
    source: str | None = None
    update_freq: str | None = None
    description: str | None = None


# ──────────── 参数 CRUD ────────────

@router.get("/parameters")
async def get_parameters():
    """获取监测参数定义列表。"""
    with get_session() as session:
        params = session.query(Parameter).order_by(Parameter.id).all()
        result = [p.to_dict() for p in params]
        return {"parameters": result, "total": len(result)}


@router.post("/parameters")
async def create_parameter(data: ParameterCreate):
    """新增监测参数。"""
    with get_session() as session:
        # 生成 ID
        max_id = session.query(Parameter).order_by(Parameter.id.desc()).first()
        next_num = (int(max_id.id[1:]) + 1) if max_id else 1
        new_id = f"P{next_num:03d}"
        key = data.name.replace(" ", "_")

        param = Parameter(
            id=new_id,
            name=data.name,
            key=key,
            unit=data.unit,
            subsystem=data.subsystem,
            location=data.location,
            normal_min=data.normal_min,
            normal_max=data.normal_max,
            source=data.source,
            update_freq=data.update_freq,
            description=data.description,
        )
        session.add(param)
        session.commit()
        return {"success": True, "id": new_id, "message": f"参数 {data.name} 已添加"}


@router.put("/parameters/{param_id}")
async def update_parameter(param_id: str, data: ParameterUpdate):
    """修改监测参数。"""
    with get_session() as session:
        param = session.query(Parameter).filter(Parameter.id == param_id).first()
        if not param:
            raise HTTPException(404, "参数不存在")
        update_data = data.model_dump(exclude_unset=True)
        for k, v in update_data.items():
            setattr(param, k, v)
        if "name" in update_data:
            param.key = update_data["name"].replace(" ", "_")
        param.updated_at = datetime.now()
        session.commit()
        return {"success": True, "message": "参数已更新"}


@router.delete("/parameters/{param_id}")
async def delete_parameter(param_id: str):
    """删除监测参数。"""
    with get_session() as session:
        param = session.query(Parameter).filter(Parameter.id == param_id).first()
        if not param:
            raise HTTPException(404, "参数不存在")
        session.delete(param)
        session.commit()
        return {"success": True, "message": "参数已删除"}


# ──────────── 基准值配置 CRUD ────────────

@router.get("/baselines")
async def get_baselines():
    """获取动态基准值模型配置。"""
    with get_session() as session:
        configs = session.query(BaselineConfig).all()
        result = [bc.to_dict() for bc in configs]
        return {"baselines": result, "total": len(result)}


# ──────────── 对标指标 CRUD ────────────

@router.get("/benchmark-indicators")
async def get_benchmark_indicators():
    """获取对标指标配置。"""
    with get_session() as session:
        indicators = session.query(BenchmarkIndicator).order_by(BenchmarkIndicator.id).all()
        result = [bi.to_dict() for bi in indicators]
        return {"indicators": result, "total": len(result)}


# ──────────── 损失项配置 CRUD ────────────

class LossVariableCreate(BaseModel):
    param_key: str
    name: str
    unit: str
    baseline: float
    best: float


class LossVariableUpdate(BaseModel):
    param_key: str | None = None
    name: str | None = None
    unit: str | None = None
    baseline: float | None = None
    best: float | None = None
    sort_order: int | None = None


@router.get("/loss-variables")
async def get_loss_variables():
    """获取所有损失项配置。"""
    with get_session() as session:
        items = session.query(LossVariableConfig).order_by(LossVariableConfig.sort_order, LossVariableConfig.id).all()
        return {"variables": [i.to_dict() for i in items], "total": len(items)}


@router.post("/loss-variables")
async def create_loss_variable(data: LossVariableCreate):
    """新增损失项配置。"""
    with get_session() as session:
        max_order = session.query(LossVariableConfig).order_by(LossVariableConfig.sort_order.desc()).first()
        next_order = (max_order.sort_order + 1) if max_order else 0
        item = LossVariableConfig(
            param_key=data.param_key,
            name=data.name,
            unit=data.unit,
            baseline=data.baseline,
            best=data.best,
            sort_order=next_order,
        )
        session.add(item)
        session.commit()
        session.refresh(item)
        return {"success": True, "id": item.id, "variable": item.to_dict()}


@router.put("/loss-variables/{item_id}")
async def update_loss_variable(item_id: int, data: LossVariableUpdate):
    """更新损失项配置。"""
    with get_session() as session:
        item = session.query(LossVariableConfig).filter(LossVariableConfig.id == item_id).first()
        if not item:
            raise HTTPException(404, "损失项不存在")
        update_data = data.model_dump(exclude_unset=True)
        for k, v in update_data.items():
            setattr(item, k, v)
        item.updated_at = datetime.now()
        session.commit()
        session.refresh(item)
        return {"success": True, "variable": item.to_dict()}


@router.delete("/loss-variables/{item_id}")
async def delete_loss_variable(item_id: int):
    """删除损失项配置。"""
    with get_session() as session:
        item = session.query(LossVariableConfig).filter(LossVariableConfig.id == item_id).first()
        if not item:
            raise HTTPException(404, "损失项不存在")
        session.delete(item)
        session.commit()
        return {"success": True, "message": "损失项已删除"}


# ──────────── XLSX 模板下载 / 导入 / 导出 ────────────

# 模板列定义
_TEMPLATE_HEADERS = ["参数名称", "单位", "子系统", "位置", "正常下限", "正常上限", "数据来源", "更新频率", "备注"]
_SUBSYSTEM_OPTIONS = ["全厂", "机组", "燃机", "压气机", "燃烧室", "透平", "发电机", "汽机", "余热锅炉", "辅助", "环境", "其他"]
_SAMPLE_ROWS = [
    ["压气机进口温度 T1", "°C", "压气机", "压气机进口", 10, 40, "时序数据库", "1s", ""],
    ["排气温度分散度", "°C", "透平", "透平出口", 0, 30, "时序数据库", "1s", "TTDM"],
]


def _build_parameter_xlsx(rows: list[list], sheet_title: str = "参数模板") -> io.BytesIO:
    """用 openpyxl 生成参数 XLSX，写入 BytesIO 返回。"""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()

    # ── 说明 sheet ──
    ws_info = wb.active
    ws_info.title = "说明"
    info_rows = [
        ["数据字典参数模板填写说明"],
        [""],
        ["字段说明："],
        ["参数名称", "必填，参数的中文标识，如 '压气机进口温度 T1'"],
        ["单位", "必填，测量单位，如 '°C', 'MPa', 'MW'"],
        ["子系统", "下拉选择：全厂/机组/燃机/压气机/燃烧室/透平/发电机/汽机/余热锅炉/辅助/环境/其他"],
        ["位置", "测点安装位置，如 '压气机进口'"],
        ["正常下限", "数值，参数正常范围下限"],
        ["正常上限", "数值，参数正常范围上限"],
        ["数据来源", "如 '时序数据库', '定期化验'，默认 '时序数据库'"],
        ["更新频率", "如 '1s', '1h', '1天'，默认 '1s'"],
        ["备注", "补充说明"],
        [""],
        ["使用步骤："],
        ["1. 在 '参数模板' sheet 中填写参数信息"],
        ["2. 参数名称和单位为必填项"],
        ["3. 子系统列有下拉校验，请从选项中选择"],
        ["4. 已存在的参数（按名称匹配）会被自动跳过"],
        ["5. 保存后上传即可批量导入"],
    ]
    for row in info_rows:
        ws_info.append(row)
    ws_info.column_dimensions["A"].width = 16
    ws_info.column_dimensions["B"].width = 60

    # ── 参数模板 sheet ──
    ws = wb.create_sheet(title=sheet_title)
    ws.append(_TEMPLATE_HEADERS)

    # 表头样式
    from openpyxl.styles import Font, PatternFill, Alignment
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")
    for col_idx, _ in enumerate(_TEMPLATE_HEADERS, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align

    # 必填标记（参数名称、单位）
    for col_idx in (1, 2):
        cell = ws.cell(row=1, column=col_idx)
        cell.value = cell.value + " *"

    # 示例数据
    for row_data in rows:
        ws.append(row_data)

    # 列宽
    col_widths = [22, 10, 12, 14, 12, 12, 14, 12, 20]
    for i, w in enumerate(col_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w

    # 子系统下拉校验（C 列，第 2 行到第 500 行）
    subsystem_list = '"' + ",".join(_SUBSYSTEM_OPTIONS) + '"'
    dv = DataValidation(type="list", formula1=subsystem_list, allow_blank=True)
    dv.error = "请从下拉列表中选择子系统"
    dv.errorTitle = "输入错误"
    ws.add_data_validation(dv)
    dv.add(f"C2:C500")

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


@router.get("/parameters/template")
async def download_parameter_template():
    """下载参数字典 XLSX 模板。"""
    buffer = _build_parameter_xlsx(_SAMPLE_ROWS)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": "attachment; filename*=UTF-8''%E6%95%B0%E6%8D%AE%E5%AD%97%E5%85%B8%E6%A8%A1%E6%9D%BF.xlsx",
        },
    )


@router.post("/parameters/import")
async def import_parameters(file: UploadFile = File(...)):
    """从 XLSX 文件批量导入参数定义。

    读取"参数模板" sheet，按名称去重后批量新增。
    返回 { success, imported, skipped, errors }。
    """
    import openpyxl

    if not file.filename or not file.filename.endswith((".xlsx", ".xls")):
        raise HTTPException(400, "请上传 .xlsx 文件")

    content = await file.read()
    if len(content) == 0:
        raise HTTPException(400, "文件为空")

    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception as e:
        raise HTTPException(400, f"文件解析失败: {e}")

    # 查找参数模板 sheet
    ws = None
    for name in wb.sheetnames:
        if "参数" in name or "模板" in name:
            ws = wb[name]
            break
    if ws is None:
        ws = wb.active

    rows = list(ws.iter_rows(min_row=2, values_only=True))

    if not rows:
        return {"success": True, "imported": 0, "skipped": 0, "errors": []}

    # 读取已有参数名称集合（去重）
    with get_session() as session:
        existing = {p.name for p in session.query(Parameter).all()}
        max_param = session.query(Parameter).order_by(Parameter.id.desc()).first()
        next_num = (int(max_param.id[1:]) + 1) if max_param else 1

        imported = 0
        skipped = 0
        errors: list[str] = []

        for row_idx, row in enumerate(rows, 2):
            # 跳过全空行
            if not row or all(v is None for v in row):
                continue

            name = str(row[0]).strip() if row[0] else ""
            unit = str(row[1]).strip() if row[1] else ""

            if not name or not unit:
                errors.append(f"第 {row_idx} 行：参数名称和单位为必填项")
                continue

            # 按名称去重
            if name in existing:
                skipped += 1
                continue

            subsystem = str(row[2]).strip() if row[2] and row[2] is not None else "其他"
            if subsystem not in _SUBSYSTEM_OPTIONS:
                subsystem = "其他"

            location = str(row[3]).strip() if row[3] and row[3] is not None else ""

            try:
                normal_min = float(row[4]) if row[4] is not None else None
            except (ValueError, TypeError):
                normal_min = None
            try:
                normal_max = float(row[5]) if row[5] is not None else None
            except (ValueError, TypeError):
                normal_max = None

            source = str(row[6]).strip() if row[6] and row[6] is not None else "时序数据库"
            update_freq = str(row[7]).strip() if row[7] and row[7] is not None else "1s"
            description = str(row[8]).strip() if len(row) > 8 and row[8] and row[8] is not None else ""

            new_id = f"P{next_num:03d}"
            key = name.replace(" ", "_")
            next_num += 1

            param = Parameter(
                id=new_id,
                name=name,
                key=key,
                unit=unit,
                subsystem=subsystem,
                location=location,
                normal_min=normal_min,
                normal_max=normal_max,
                source=source,
                update_freq=update_freq,
                description=description,
            )
            session.add(param)
            existing.add(name)
            imported += 1

        session.commit()

    return {
        "success": True,
        "imported": imported,
        "skipped": skipped,
        "errors": errors,
    }


@router.get("/parameters/export")
async def export_parameters():
    """导出当前所有参数定义为 XLSX 文件。"""
    with get_session() as session:
        params = session.query(Parameter).order_by(Parameter.id).all()
        rows = [
            [
                p.name,
                p.unit,
                p.subsystem or "",
                p.location or "",
                p.normal_min,
                p.normal_max,
                p.source or "",
                p.update_freq or "",
                p.description or "",
            ]
            for p in params
        ]

    buffer = _build_parameter_xlsx(rows, sheet_title="参数数据")
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={
            "Content-Disposition": "attachment; filename*=UTF-8''%E6%95%B0%E6%8D%AE%E5%AD%97%E5%85%B8%E5%AF%BC%E5%87%BA.xlsx",
        },
    )


# ──────────── 瀑布图配置 ────────────

class WaterfallConfigUpdate(BaseModel):
    total_key: str = ""
    subsystem_keys: list[str] = []


@router.get("/waterfall-config")
async def get_waterfall_config():
    """获取瀑布图全局配置。"""
    with get_session() as session:
        config = session.query(WaterfallConfig).first()
        if not config:
            return {"total_key": "", "subsystem_keys": []}
        return config.to_dict()


@router.put("/waterfall-config")
async def save_waterfall_config(data: WaterfallConfigUpdate):
    """保存瀑布图全局配置。"""
    with get_session() as session:
        config = session.query(WaterfallConfig).first()
        if not config:
            config = WaterfallConfig(
                total_key=data.total_key,
                subsystem_keys=json.dumps(data.subsystem_keys, ensure_ascii=False),
            )
            session.add(config)
        else:
            config.total_key = data.total_key
            config.subsystem_keys = json.dumps(data.subsystem_keys, ensure_ascii=False)
            config.updated_at = datetime.now()
        session.commit()
        session.refresh(config)
        return {"success": True, "config": config.to_dict()}
